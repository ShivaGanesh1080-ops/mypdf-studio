import os
import sys
import requests
import pymupdf as fitz
from PIL import Image

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

BASE_URL = "http://127.0.0.1:8000"

def test_new_routes():
    print("=== Testing New 'Create PDF' & Office Conversion Routes ===\n")
    test_dir = os.path.join(os.path.dirname(__file__), "test_new_tmp")
    os.makedirs(test_dir, exist_ok=True)
    
    # 1. Images (JPG + PNG) to PDF with page_size and margin
    img1 = os.path.join(test_dir, "pic1.jpg")
    img2 = os.path.join(test_dir, "pic2.png")
    Image.new("RGB", (400, 300), color="orange").save(img1, "JPEG")
    Image.new("RGB", (300, 500), color="purple").save(img2, "PNG")
    
    with open(img1, "rb") as f1, open(img2, "rb") as f2:
        files = [("files", ("pic1.jpg", f1, "image/jpeg")), ("files", ("pic2.png", f2, "image/png"))]
        data = {"page_size": "a4", "orientation": "auto", "margin": 15}
        res = requests.post(f"{BASE_URL}/api/convert/images-to-pdf", files=files, data=data)
        assert res.status_code == 200, f"Images to PDF failed: {res.text}"
        doc = fitz.open(stream=res.content, filetype="pdf")
        assert len(doc) == 2, f"Expected 2 pages, got {len(doc)}"
        doc.close()
    print("[PASS] 1. /api/convert/images-to-pdf (JPG + PNG to A4 PDF with 15pt margins)")

    # 2. Plain Text to PDF
    txt_path = os.path.join(test_dir, "notes.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("MyPDF Studio Report\n\nThis is a sample text file converted to PDF.\nLine 1: Success\nLine 2: 100% Private")
    with open(txt_path, "rb") as f:
        res = requests.post(f"{BASE_URL}/api/convert/text-to-pdf", files={"file": ("notes.txt", f, "text/plain")})
        assert res.status_code == 200, f"Text to PDF failed: {res.text}"
        doc = fitz.open(stream=res.content, filetype="pdf")
        assert len(doc) >= 1, "Expected at least 1 page"
        doc.close()
    print("[PASS] 2. /api/convert/text-to-pdf (Plain text converted to clean PDF)")

    # 3. Markdown to PDF
    md_path = os.path.join(test_dir, "readme.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("# Project Overview\n\n## Key Features\n- 100% Private\n- No Limits\n- AI Powered\n")
    with open(md_path, "rb") as f:
        res = requests.post(f"{BASE_URL}/api/convert/md-to-pdf", files={"file": ("readme.md", f, "text/markdown")})
        assert res.status_code == 200, f"Markdown to PDF failed: {res.text}"
        doc = fitz.open(stream=res.content, filetype="pdf")
        assert len(doc) >= 1, "Expected at least 1 page"
        doc.close()
    print("[PASS] 3. /api/convert/md-to-pdf (Markdown converted to styled PDF)")

    # 4. Generate a sample Word .docx and convert to PDF
    import docx
    docx_path = os.path.join(test_dir, "sample.docx")
    w_doc = docx.Document()
    w_doc.add_heading("Automated Word Document", 0)
    w_doc.add_paragraph("This document was generated and converted to PDF using MyPDF Studio.")
    w_doc.save(docx_path)
    
    with open(docx_path, "rb") as f:
        res = requests.post(f"{BASE_URL}/api/convert/docx-to-pdf", files={"file": ("sample.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert res.status_code == 200, f"Word to PDF failed: {res.text}"
        doc = fitz.open(stream=res.content, filetype="pdf")
        assert len(doc) >= 1, "Expected at least 1 page"
        doc.close()
    print("[PASS] 4. /api/convert/docx-to-pdf (Word .docx converted to PDF)")

    # 5. Generate a sample PowerPoint .pptx and convert to PDF
    from pptx import Presentation
    pptx_path = os.path.join(test_dir, "sample.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "MyPDF Presentation"
    slide.placeholders[1].text = "Converted via PowerPoint Engine"
    prs.save(pptx_path)

    with open(pptx_path, "rb") as f:
        res = requests.post(f"{BASE_URL}/api/convert/pptx-to-pdf", files={"file": ("sample.pptx", f, "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
        assert res.status_code == 200, f"PPTX to PDF failed: {res.text}"
        doc = fitz.open(stream=res.content, filetype="pdf")
        assert len(doc) >= 1, "Expected at least 1 page"
        doc.close()
    print("[PASS] 5. /api/convert/pptx-to-pdf (PowerPoint .pptx converted to PDF)")

    # 6. Generate a sample Excel .xlsx and convert to PDF
    import openpyxl
    xlsx_path = os.path.join(test_dir, "sample.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Metric"
    ws["B1"] = "Value"
    ws["A2"] = "Active Users"
    ws["B2"] = 12500
    wb.save(xlsx_path)

    with open(xlsx_path, "rb") as f:
        res = requests.post(f"{BASE_URL}/api/convert/excel-to-pdf", files={"file": ("sample.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert res.status_code == 200, f"Excel to PDF failed: {res.text}"
        doc = fitz.open(stream=res.content, filetype="pdf")
        assert len(doc) >= 1, "Expected at least 1 page"
        doc.close()
    print("[PASS] 6. /api/convert/excel-to-pdf (Excel .xlsx converted to PDF)")

    # Cleanup
    import shutil
    shutil.rmtree(test_dir)
    print("\nALL 'CREATE PDF' & OFFICE CONVERTERS TESTED & VERIFIED SUCCESSFULLY!")

if __name__ == "__main__":
    test_new_routes()
